
from __future__ import annotations
from io import BytesIO
from docx import Document
from docx.shared import Pt


def build_report(model_info, opt, kkt):
    doc = Document()
    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10)

    doc.add_heading("ĐỚP TOOL — THUYẾT MINH TỐI ƯU KẾT CẤU THÉP", 0)
    doc.add_paragraph(
        "Quy trình: mô hình hóa khung → phân tích FEM → xác định nội lực/chuyển vị "
        "→ ứng viên tối ưu bằng biến phân → tối ưu số có ràng buộc chuyển vị → kiểm tra KKT."
    )

    doc.add_heading("1. Mô hình và phân tích FEM", level=1)
    doc.add_paragraph(
        f"Số nút: {model_info['n_nodes']}; số phần tử: {model_info['n_elements']}; "
        f"số gối: {model_info['n_supports']}. "
        f"Chuyển vị lớn nhất trước tối ưu: {opt['d_initial']:.6g} m."
    )

    doc.add_heading("2. Hàm mục tiêu", level=1)
    doc.add_paragraph(
        "Tối thiểu hóa thể tích thép: J = Σ A_e L_e. "
        "Với mật độ thép không đổi, tối thiểu thể tích tương đương tối thiểu khối lượng."
    )

    doc.add_heading("3. Biến phân và Euler–Lagrange", level=1)
    doc.add_paragraph(
        "Đối với trường diện tích liên tục A(x), xét J[A] = ∫A(x)dx và ràng buộc "
        "năng lượng biến dạng dọc ∫N(x)^2/[E A(x)]dx ≤ C. "
        "Vì hàm dưới dấu tích phân không chứa A'(x), phương trình Euler–Lagrange "
        "trở thành điều kiện đại số ∂F/∂A = 0."
    )
    doc.add_paragraph(
        "Điều kiện KKT liên tục cho nghiệm dọc trục cho dạng: "
        "A*(x) = max(|N(x)|/σ_allow, |N(x)|√(λ/E)). "
        "λ được chọn sao cho ràng buộc chuyển vị/năng lượng hoạt động khi cần."
    )
    doc.add_paragraph(
        "Đối với uốn, FEM cung cấp M(x); nguyên lý tương tự cho thành phần độ cứng "
        "uốn sử dụng ∫M(x)^2/[E I(x)]dx. Trong phiên bản này ứng viên biến phân được "
        "dùng để khởi tạo cho bài toán tối ưu hữu hạn chiều có kiểm tra chuyển vị FEM trực tiếp."
    )

    doc.add_heading("4. Kết quả tối ưu", level=1)
    doc.add_paragraph(
        f"Thể tích tương đối: {opt['objective_initial']:.6g} → {opt['objective_final']:.6g}. "
        f"Chuyển vị lớn nhất: {opt['d_initial']:.6g} → {opt['d_final']:.6g} m. "
        f"Trạng thái solver: {opt['message']}."
    )
    table = doc.add_table(rows=1, cols=5)
    hdr = table.rows[0].cells
    for c, t in zip(hdr, ["Element", "L (m)", "A trước (m²)", "A tối ưu (m²)", "I tối ưu (m⁴)"]):
        c.text = t
    for i, (L, a0, a1, ii) in enumerate(zip(opt["lengths"], opt["A_initial"], opt["A_final"], opt["I_final"])):
        row = table.add_row().cells
        for c, t in zip(row, [str(i), f"{L:.4f}", f"{a0:.6g}", f"{a1:.6g}", f"{ii:.6g}"]):
            c.text = t

    doc.add_heading("5. Kiểm tra KKT", level=1)
    doc.add_paragraph(
        f"Primal feasibility: max(g) = {kkt['primal_max_violation']:.3e}. "
        f"Stationarity norm = {kkt['stationarity_norm']:.3e}. "
        f"Dual feasibility = {kkt['dual_feasible']}. "
        f"Complementarity residual = {kkt['complementarity_residual']:.3e}."
    )
    doc.add_paragraph(
        "Kết luận KKT được xem là đạt khi vi phạm primal gần 0, gradient stationarity nhỏ, "
        "các multiplier không âm và tích λᵢgᵢ gần 0. Đây là kiểm tra số tại nghiệm thu được, "
        "không phải chứng minh toàn cục nếu bài toán không lồi."
    )

    doc.add_heading("6. Kết luận", level=1)
    doc.add_paragraph(
        "Tiết diện tối ưu phải được phân tích lại bằng FEM với chính A và I sau tối ưu. "
        "Chỉ khi đồng thời thỏa mãn chuyển vị, điều kiện độ bền và các ràng buộc hình học "
        "mới được xem là phương án thiết kế hợp lệ. Các kiểm tra ổn định, cấu tạo, liên kết "
        "và lựa chọn tiết diện thương mại cần thực hiện ở bước thiết kế thép tiếp theo."
    )

    bio = BytesIO()
    doc.save(bio)
    return bio.getvalue()
